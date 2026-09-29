import re
from dataclasses import dataclass
from typing import Any
from openpyxl import load_workbook

HEADER_ALIASES = {
    "author_id": {"author id", "author uid", "author_uid", "authorid"},
    "book_id": {"book id", "book_id", "novel bookid"},
    "show_id": {"show id", "show_id"},
    "status": {"payout status", "payment status", "status", "payout_status"},
    "utr": {"utr", "utr no.", "utr no", "utr number"},
    "account_number": {"user bank account", "account number", "beneficiary account no.", "account_no"},
    "ifsc": {"ifsc code", "ifsc_code", "ifsc"},
    "bank_name": {"bank name", "bank_name"},
    "payout_id": {"payout id", "payout_id"},
    "amount_before_tax": {"amount before tax", "requested gross", "amount", "gross"},
    "amount_after_tax": {"amount after tax", "processed net", "requested net", "net amount"},
    "tds": {"tds%", "tds", "processed tds", "tds amount"},
    "transaction_date": {"transaction date", "value date"},
    "payment_period": {"payout month", "payment month"},
}

KIND_RULES = [
    ("uwt", ("uwt format", "uwt final", "uwt")),
    ("finance_success", ("fin_success", "finance success")),
    ("finance_reversal", ("fin_reversal", "reversal")),
    ("payment_status", ("payment status dump", "payment status")),
    ("show_book_mapping", ("show book mapping", "series ret")),
    ("author_level", ("author level",)),
    ("incentive", ("inc final", "inc dump")),
    ("revenue_share", ("rs final", "rs dump", "revenue share dump")),
    ("recovery", ("recovery",)),
    ("pan_status", ("pan status", "pan_status")),
    ("bank_details", ("bank details", "bank_names")),
]

def norm(v: Any) -> str:
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).strip().lower().replace("\n", " "))

def sheet_kind(title: str) -> str:
    n = norm(title)
    for kind, needles in KIND_RULES:
        if any(x in n for x in needles):
            return kind
    return "other"

def find_header_row(ws, scan_rows: int = 15) -> tuple[int | None, list[str]]:
    best = None
    for i, row in enumerate(
        ws.iter_rows(min_row=1, max_row=min(scan_rows, ws.max_row or scan_rows), values_only=True), 1
    ):
        vals = [norm(v) for v in row]
        hits = 0
        for v in vals:
            for aliases in HEADER_ALIASES.values():
                if any(v == a or a in v for a in aliases):
                    hits += 1
                    break
        nonempty = sum(bool(v) for v in vals)
        if nonempty and (best is None or hits > best[0]):
            best = (hits, i, vals)
    if not best:
        return None, []
    return best[1], best[2]

def header_map(values: list[str]) -> dict[str, int]:
    result = {}
    for i, v in enumerate(values):
        if not v:
            continue
        for key, aliases in HEADER_ALIASES.items():
            if key in result:
                continue
            if any(v == a or a in v for a in aliases):
                result[key] = i
    return result

def parse_bank_json(value: Any) -> dict[str, str | None]:
    import json
    if not isinstance(value, str):
        return {"account_number": None, "ifsc": None, "bank_name": None}
    try:
        data = json.loads(value)
        return {
            "account_number": str(data.get("account_number")) if data.get("account_number") is not None else None,
            "ifsc": str(data.get("ifsc")) if data.get("ifsc") is not None else None,
            "bank_name": str(data.get("bank_name")) if data.get("bank_name") is not None else None,
        }
    except Exception:
        return {"account_number": None, "ifsc": None, "bank_name": None}

def inspect_workbook(path: str, sample_rows: int = 3) -> dict[str, Any]:
    wb = load_workbook(path, read_only=True, data_only=False)
    summaries = []
    for ws in wb.worksheets:
        header_row, headers = find_header_row(ws)
        mapping = header_map(headers) if header_row else {}
        row_count = 0
        samples = []
        if header_row:
            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if any(v not in (None, "") for v in row):
                    row_count += 1
                    if len(samples) < sample_rows:
                        samples.append([str(v)[:120] if v is not None else None for v in row])
        summaries.append({
            "name": ws.title,
            "kind": sheet_kind(ws.title),
            "header_row": header_row,
            "row_count": row_count,
            "mapped_fields": sorted(mapping.keys()),
            "sample_rows": samples,
        })
    return {"filename": path.rsplit("/", 1)[-1], "sheets": summaries}
