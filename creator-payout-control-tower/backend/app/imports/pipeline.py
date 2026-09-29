import hashlib
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from .inspector import find_header_row, header_map, sheet_kind, parse_bank_json
from ..services.status import normalize_payment_status

def sha256_file(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def infer_period(filename: str) -> str | None:
    m = re.search(r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[\' ]?(\d{2})", filename, re.I)
    if not m:
        return None
    return f"{m.group(1).title()} 20{m.group(2)}"

def jsonable(v: Any) -> Any:
    if isinstance(v, datetime):
        return v.isoformat()
    if isinstance(v, (str, int, float, bool)) or v is None:
        return v
    return str(v)

def row_to_dict(headers: list[str], row: list[Any]) -> dict[str, Any]:
    out = {}
    for i, value in enumerate(row):
        key = headers[i].strip() if i < len(headers) and headers[i] else f"column_{i+1}"
        if key in out:
            key = f"{key}_{i+1}"
        out[key] = jsonable(value)
    return out

def normalize_row(kind: str, mapping: dict[str, int], row: list[Any]) -> dict[str, Any]:
    def val(key: str):
        i = mapping.get(key)
        return row[i] if i is not None and i < len(row) else None

    raw_bank = val("account_number")
    bank = parse_bank_json(raw_bank)
    status_raw = val("status")
    status = normalize_payment_status(str(status_raw) if status_raw not in (None, "") else None).value

    return {
        "kind": kind,
        "author_id": str(val("author_id")).strip() if val("author_id") not in (None, "") else None,
        "book_id": str(val("book_id")).strip() if val("book_id") not in (None, "") else None,
        "show_id": str(val("show_id")).strip() if val("show_id") not in (None, "") else None,
        "status_raw": str(status_raw).strip() if status_raw not in (None, "") else None,
        "status": status,
        "utr": str(val("utr")).strip() if val("utr") not in (None, "") else None,
        "account_number": bank["account_number"] or (str(raw_bank).strip() if raw_bank not in (None, "") and kind != "uwt" else None),
        "ifsc": bank["ifsc"] or (str(val("ifsc")).strip() if val("ifsc") not in (None, "") else None),
        "bank_name": bank["bank_name"] or (str(val("bank_name")).strip() if val("bank_name") not in (None, "") else None),
        "payout_id": str(val("payout_id")).strip() if val("payout_id") not in (None, "") else None,
        "amount_before_tax": val("amount_before_tax"),
        "amount_after_tax": val("amount_after_tax"),
        "tds": val("tds"),
        "transaction_date": jsonable(val("transaction_date")) if val("transaction_date") is not None else None,
        "payment_period": str(val("payment_period")).strip() if val("payment_period") not in (None, "") else None,
    }

def ingest_workbook(path: str) -> dict[str, Any]:
    p = Path(path)
    wb = load_workbook(path, read_only=True, data_only=False)
    batch_id = sha256_file(path)
    period = infer_period(p.name)
    total_rows = 0
    events = []
    sheets = []

    for ws in wb.worksheets:
        header_row, headers = find_header_row(ws)
        if not header_row:
            continue
        mapping = header_map(headers)
        kind = sheet_kind(ws.title)
        count = 0
        for row_number, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
            if not any(v not in (None, "") for v in row):
                continue
            count += 1
            total_rows += 1
            if kind in {"uwt", "finance_success", "finance_reversal", "payment_status"}:
                item = normalize_row(kind, mapping, list(row))
                item.update({
                    "source_file": p.name,
                    "source_sheet": ws.title,
                    "source_row": row_number,
                    "inferred_period": period,
                    "raw": row_to_dict(headers, list(row)),
                })
                events.append(item)
        sheets.append({"name": ws.title, "kind": kind, "header_row": header_row, "row_count": count})

    return {
        "batch_id": batch_id,
        "filename": p.name,
        "inferred_period": period,
        "total_rows": total_rows,
        "payment_event_rows": len(events),
        "sheets": sheets,
        "payment_events": events,
    }
