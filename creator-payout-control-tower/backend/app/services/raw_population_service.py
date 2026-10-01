from collections import defaultdict
from decimal import Decimal
from pathlib import Path
import csv
import json
import uuid

from openpyxl import load_workbook

from ..db.session import SessionLocal
from ..models import Earning, MonthlyRun, RawImportRow


def s(v):
    return str(v).strip() if v not in (None, "") else ""


def d(v):
    try:
        return Decimal(str(v).replace(",", ""))
    except Exception:
        return Decimal("0")


def norm(v):
    return "".join(ch for ch in s(v).lower() if ch.isalnum())


def find_col(headers, names):
    hs = [norm(h) for h in headers]
    wanted = [norm(x) for x in names]
    for w in wanted:
        for i, h in enumerate(hs):
            if h == w:
                return i
    for w in wanted:
        if not w:
            continue
        for i, h in enumerate(hs):
            if w in h or h in w:
                return i
    return None


def detect_header(rows):
    marker_groups = [
        ("bookid", "book"),
        ("authorid", "authoruid", "author"),
        ("showid", "show"),
        ("iptype", "contenttype", "content"),
        ("pay", "status"),
    ]
    for index, raw in enumerate(rows[:50]):
        headers = [s(v) for v in raw]
        score = 0
        for group in marker_groups:
            if any(any(term in norm(h) or norm(h) in term for term in group) for h in headers):
                score += 1
        if score >= 2:
            return index, headers
    return None, None


def read_source(path, kind):
    path_obj = Path(path)
    if path_obj.suffix.lower() == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            sheets = [(path_obj.name, list(csv.reader(handle)))]
    else:
        wb = load_workbook(path, read_only=True, data_only=True)
        sheets = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]

    rows = []
    errors = []

    for sheet_name, matrix in sheets:
        header_index, headers = detect_header(matrix)
        if headers is None:
            if any(any(s(v) for v in row) for row in matrix):
                errors.append({
                    "file": path_obj.name,
                    "sheet": sheet_name,
                    "reason": "Could not detect a usable header row containing source identity columns.",
                })
            continue

        for source_row, raw in enumerate(matrix[header_index + 1:], header_index + 2):
            if not any(s(v) for v in raw):
                continue

            def get(*names):
                index = find_col(headers, names)
                return raw[index] if index is not None and index < len(raw) else None

            author_id = s(get("Author ID", "Author UID", "AuthorID", "Author_uid"))
            book_id = s(get("Book ID", "Book_id", "Novel bookID", "BookID"))
            show_id = s(get("Show ID", "Show_id", "Audio Show_ID", "PGC Show ID", "ShowID"))
            ip_type = s(get("IP Type", "Content Type", "Content_Type", "Content"))
            contract_type = s(get("Contract Type", "Contract_Type"))
            pay_flag = s(get("Pay?", "Pay", "Eligibility", "Eligible"))
            status = s(get("Status", "Payment Status"))

            marker = " ".join([ip_type, contract_type, pay_flag, status]).upper()
            if "N2A" in marker or "A2A" in marker:
                content_type = "N2A/A2A"
            elif "SERIES" in marker:
                content_type = "SERIES"
            else:
                content_type = "NOVEL"

            gross_value = get("Gross", "Gross Amount", "Amount", "Balance amount", "Total Revenue")
            tds_value = get("TDS Amount", "TDS", "Processed TDS")
            net_value = get("Net Amount", "Processed Net", "Net")

            identity_present = bool(author_id or book_id or show_id)
            if not identity_present:
                errors.append({
                    "file": path_obj.name,
                    "sheet": sheet_name,
                    "row": source_row,
                    "reason": "Row has no Author ID, Book ID, or Show ID. Row retained in raw storage but cannot appear in an Author-level view.",
                })

            raw_payload = {}
            for i in range(max(len(headers), len(raw))):
                key = headers[i] if i < len(headers) and headers[i] else f"column_{i + 1}"
                raw_payload[key] = raw[i] if i < len(raw) else None

            rows.append({
                "source_kind": kind,
                "source_file": path_obj.name,
                "source_sheet": sheet_name,
                "source_row": source_row,
                "author_id": author_id or None,
                "book_id": book_id or None,
                "show_id": show_id or None,
                "content_type": content_type,
                "pay_flag": pay_flag or None,
                "status_raw": status or None,
                "gross": d(gross_value),
                "tds": d(tds_value),
                "net": d(net_value),
                "raw": raw_payload,
            })

    return rows, errors


def aggregate(rows):
    grouped = defaultdict(lambda: {
        "inc": Decimal("0"),
        "rs": Decimal("0"),
        "gross": Decimal("0"),
        "book_ids": set(),
        "show_ids": set(),
        "content_types": set(),
        "references": [],
        "row_count": 0,
        "novel_rows": 0,
        "series_rows": 0,
        "n2a_rows": 0,
    })

    for row in rows:
        author_id = s(row.get("author_id"))
        if not author_id:
            continue
        group = grouped[author_id]
        amount = row.get("gross") or Decimal("0")

        if row["source_kind"] == "INCENTIVE_DUMP":
            group["inc"] += amount
        else:
            group["rs"] += amount

        group["gross"] += amount
        if row.get("book_id"):
            group["book_ids"].add(row["book_id"])
        if row.get("show_id"):
            group["show_ids"].add(row["show_id"])

        group["content_types"].add(row.get("content_type") or "NOVEL")
        group["references"].append(
            f'{row["source_kind"]}:{row["source_file"]}:{row["source_sheet"]}:R{row["source_row"]}'
        )
        group["row_count"] += 1

        if row.get("content_type") == "SERIES":
            group["series_rows"] += 1
        elif row.get("content_type") == "N2A/A2A":
            group["n2a_rows"] += 1
        else:
            group["novel_rows"] += 1

    result = []
    for author_id, group in sorted(grouped.items()):
        result.append({
            "author_id": author_id,
            "incentive_gross": float(group["inc"]),
            "revenue_share_gross": float(group["rs"]),
            "gross": float(group["gross"]),
            "book_ids": ", ".join(sorted(group["book_ids"])),
            "show_ids": ", ".join(sorted(group["show_ids"])),
            "book_count": len(group["book_ids"]),
            "show_count": len(group["show_ids"]),
            "content_types": " / ".join(sorted(group["content_types"])),
            "source_row_count": group["row_count"],
            "novel_rows": group["novel_rows"],
            "series_rows": group["series_rows"],
            "n2a_rows": group["n2a_rows"],
            "source_references": " | ".join(group["references"]),
        })

    return result


def aggregate_type(rows, content_type):
    return aggregate([row for row in rows if row.get("content_type") == content_type])


def populate_two_dumps_raw(inc_path, rs_path, period):
    db = SessionLocal()
    try:
        inc_rows, inc_errors = read_source(inc_path, "INCENTIVE_DUMP")
        rs_rows, rs_errors = read_source(rs_path, "REVENUE_SHARE_DUMP")
        all_rows = inc_rows + rs_rows
        errors = inc_errors + rs_errors

        run_id = uuid.uuid4().hex
        db.query(Earning).filter(Earning.payment_period == period).delete(synchronize_session=False)
        db.query(MonthlyRun).filter(MonthlyRun.period == period).delete(synchronize_session=False)

        author_level = aggregate(all_rows)
        novel_level = aggregate_type(all_rows, "NOVEL")
        series_level = aggregate_type(all_rows, "SERIES")
        n2a_level = aggregate_type(all_rows, "N2A/A2A")

        run = MonthlyRun(
            id=run_id,
            period=period,
            incentive_file=Path(inc_path).name,
            revenue_share_file=Path(rs_path).name,
            status="POPULATED",
            input_rows=len(all_rows),
            output_author_rows=len(author_level),
            qc_blocked=0,
            qc_review=0,
        )
        db.add(run)

        for row in all_rows:
            db.add(Earning(
                id=uuid.uuid4().hex,
                payment_period=period,
                reward_type="INCENTIVE" if row["source_kind"] == "INCENTIVE_DUMP" else "REVENUE_SHARE",
                author_id=row["author_id"],
                book_id=row["book_id"],
                show_id=row["show_id"],
                content_type=row["content_type"],
                pay_flag=row["pay_flag"],
                exclusion_reason=None,
                gross=row["gross"],
                tds=row["tds"],
                net=row["net"],
                source_file=row["source_file"],
                source_sheet=row["source_sheet"],
                source_row=row["source_row"],
            ))
            db.add(RawImportRow(
                id=uuid.uuid4().hex,
                batch_id=run_id,
                source_sheet=row["source_sheet"],
                source_row=row["source_row"],
                sheet_kind=row["source_kind"],
                payload=json.dumps(row["raw"], default=str),
            ))

        db.commit()

        return {
            "period": period,
            "run_id": run_id,
            "status": "POPULATED",
            "stage": "SEGREGATED",
            "input_summary": {
                "incentive_rows": len(inc_rows),
                "revenue_share_rows": len(rs_rows),
                "total_source_rows": len(all_rows),
            },
            "segregation_summary": {
                "total_rows": len(all_rows),
                "distinct_authors": len(author_level),
                "novel_rows": sum(row["novel_rows"] for row in author_level),
                "novel_authors": len(novel_level),
                "series_rows": sum(row["series_rows"] for row in author_level),
                "series_authors": len(series_level),
                "n2a_rows": sum(row["n2a_rows"] for row in author_level),
                "n2a_authors": len(n2a_level),
                "rows_with_identity": sum(1 for row in all_rows if row["author_id"] or row["book_id"] or row["show_id"]),
                "rows_without_identity": sum(1 for row in all_rows if not (row["author_id"] or row["book_id"] or row["show_id"])),
            },
            "errors": errors,
            "author_level_rows": author_level,
            "author_level_novel_rows": novel_level,
            "author_level_series_rows": series_level,
            "author_level_n2a_rows": n2a_level,
            "message": "Raw monthly dumps were read and segregated. No eligibility filtering, compliance filtering, or payout calculation was applied.",
        }
    finally:
        db.close()
