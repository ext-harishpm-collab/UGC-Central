import uuid
from decimal import Decimal
from openpyxl import load_workbook
from ..db.session import SessionLocal
from ..models import Earning,ContentItem

def s(v):return str(v).strip() if v not in (None,"") else None
def num(v):
    try:return Decimal(str(v))
    except Exception:return None
def ix(headers,*terms):
    for i,h in enumerate(headers):
        x=(h or "").lower()
        if any(t in x for t in terms):return i
    return None

def ingest_earnings(path,period=None):
    wb=load_workbook(path,read_only=True,data_only=True);db=SessionLocal()
    added=excluded=0
    try:
        mappings={}
        for sn in wb.sheetnames:
            if sn.lower()=="show book mapping":
                ws=wb[sn];hs=[s(v) or "" for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
                ib=ix(hs,"book id");isid=ix(hs,"show id");ia=ix(hs,"author uid")
                if ib is not None:
                    for r in ws.iter_rows(min_row=2,values_only=True):
                        b=s(r[ib]);mappings[b]=(s(r[isid]) if isid is not None else None,s(r[ia]) if ia is not None else None)
        for sn in wb.sheetnames:
            low=sn.lower()
            if "inc final" not in low and "rs final" not in low:continue
            reward="INCENTIVE" if "inc final" in low else "REVENUE_SHARE"
            ws=wb[sn];hs=[s(v) or "" for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
            ip=ix(hs,"ip type");pay=ix(hs,"pay?","pay");book=ix(hs,"book id");gross=ix(hs,"gross");tds=ix(hs,"tds amount","tds");net=ix(hs,"net amount","net");rs=ix(hs,"revenue%","revenue %")
            for rn,r in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
                bid=s(r[book]) if book is not None else None
                if not bid:continue
                ctype=s(r[ip]) if ip is not None else None
                ctype=("N2A/A2A" if ctype and ("N2A" in ctype.upper() or "A2A" in ctype.upper()) else "SERIES" if ctype and "SERIES" in ctype.upper() else "NOVEL" if ctype and "NOVEL" in ctype.upper() else ctype)
                show,author=mappings.get(bid,(None,None))
                pf=s(r[pay]) if pay is not None else None
                reason=None
                if reward=="INCENTIVE" and ctype=="N2A/A2A":reason="N2A/A2A incentive exclusion"
                if pf and pf.lower() not in {"yes","paid","eligible"} and not reason:reason=f"Source Pay flag: {pf}"
                db.add(Earning(id=uuid.uuid4().hex,payment_period=period,reward_type=reward,author_id=author,book_id=bid,show_id=show,content_type=ctype,pay_flag=pf,exclusion_reason=reason,gross=num(r[gross]) if gross is not None else None,tds=num(r[tds]) if tds is not None else None,net=num(r[net]) if net is not None else None,product_rs=num(r[rs]) if rs is not None else None,source_file=path.rsplit("/",1)[-1],source_sheet=sn,source_row=rn))
                added+=1
                excluded+=1 if reason else 0
        db.commit();return {"earnings_added":added,"excluded":excluded}
    except Exception:db.rollback();raise
    finally:db.close()

def preview_author_payouts(period=None):
    db=SessionLocal()
    try:
        rows=db.query(Earning).filter(Earning.exclusion_reason.is_(None))
        if period:rows=rows.filter(Earning.payment_period==period)
        agg={}
        for e in rows.all():
            if not e.author_id:continue
            a=agg.setdefault(e.author_id,{"author_id":e.author_id,"gross_incentive":Decimal("0"),"gross_revenue_share":Decimal("0"),"other_earnings":Decimal("0"),"tds":Decimal("0"),"earning_count":0})
            if e.reward_type=="INCENTIVE":a["gross_incentive"]+=e.gross or 0
            elif e.reward_type=="REVENUE_SHARE":a["gross_revenue_share"]+=e.gross or 0
            else:a["other_earnings"]+=e.gross or 0
            a["tds"]+=e.tds or 0;a["earning_count"]+=1
        return [{k:(float(v) if isinstance(v,Decimal) else v) for k,v in x.items()} for x in agg.values()]
    finally:db.close()
