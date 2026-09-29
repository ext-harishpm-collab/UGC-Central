import uuid
from decimal import Decimal
from openpyxl import load_workbook
from ..db.session import SessionLocal
from ..models import ContractRecord,PanRecord,ComplianceRecord

def s(v):return str(v).strip() if v not in (None,"") else None
def d(v):
    try:return Decimal(str(v))
    except:return None
def ix(hs,*terms):
    for i,h in enumerate(hs):
        x=(h or "").lower()
        if any(t.lower() in x for t in terms):return i
    return None

def import_reference_workbook(path,kind):
    wb=load_workbook(path,read_only=True,data_only=True);db=SessionLocal();out={"kind":kind,"rows":0,"loaded":0,"review":0}
    try:
        for ws in wb.worksheets:
            hs=[s(v) or "" for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
            ia=ix(hs,"author id","author uid","authorid");ib=ix(hs,"book id","book_id");isid=ix(hs,"show id","show_id")
            if kind=="contract":
                irs=ix(hs,"contractual rs","contract rs","revenue share %","revenue share");ict=ix(hs,"contract type","contract")
            elif kind=="pan":
                ip=ix(hs,"pan","pan no","pan number");ist=ix(hs,"pan status","status");itr=ix(hs,"tds %","tds rate","tds percentage")
            else:
                ist=ix(hs,"compliance status","status");ireason=ix(hs,"reason","remarks","comment");idate=ix(hs,"effective date","date")
            for rn,r in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
                if not any(v not in (None,"") for v in r):continue
                out["rows"]+=1
                aid=s(r[ia]) if ia is not None and ia<len(r) else None
                bid=s(r[ib]) if ib is not None and ib<len(r) else None
                sid=s(r[isid]) if isid is not None and isid<len(r) else None
                if kind=="contract":
                    rs=d(r[irs]) if irs is not None and irs<len(r) else None
                    if not any([aid,bid,sid]) or rs is None:out["review"]+=1;continue
                    db.add(ContractRecord(id=uuid.uuid4().hex,author_id=aid,book_id=bid,show_id=sid,contract_type=s(r[ict]) if ict is not None and ict<len(r) else None,contract_rs=rs,source_file=path.rsplit("/",1)[-1],source_sheet=ws.title,source_row=rn));out["loaded"]+=1
                elif kind=="pan":
                    if not aid:out["review"]+=1;continue
                    rate=d(r[itr]) if itr is not None and itr<len(r) else None
                    if rate is not None and rate>1:rate=rate/100
                    db.add(PanRecord(id=uuid.uuid4().hex,author_id=aid,pan=s(r[ip]) if ip is not None and ip<len(r) else None,pan_status=s(r[ist]) if ist is not None and ist<len(r) else None,validated_tds_rate=rate,source_file=path.rsplit("/",1)[-1],source_sheet=ws.title,source_row=rn));out["loaded"]+=1
                else:
                    if not any([aid,bid,sid]):out["review"]+=1;continue
                    db.add(ComplianceRecord(id=uuid.uuid4().hex,author_id=aid,book_id=bid,show_id=sid,status=s(r[ist]) if ist is not None and ist<len(r) else None,
                                             reason=s(r[ireason]) if ireason is not None and ireason<len(r) else None,effective_date=r[idate] if idate is not None and idate<len(r) else None,
                                             source_file=path.rsplit("/",1)[-1],source_sheet=ws.title,source_row=rn));out["loaded"]+=1
        db.commit();return out
    except Exception:db.rollback();raise
    finally:db.close()

def author_controls(author_id,period=None):
    db=SessionLocal()
    try:
        contracts=db.query(ContractRecord).filter(ContractRecord.author_id==author_id).order_by(ContractRecord.effective_from.desc().nullslast()).all()
        pans=db.query(PanRecord).filter(PanRecord.author_id==author_id).order_by(PanRecord.validated_at.desc().nullslast()).all()
        comps=db.query(ComplianceRecord).filter(ComplianceRecord.author_id==author_id).all()
        return {"contracts":[{"book_id":x.book_id,"show_id":x.show_id,"type":x.contract_type,"rs":float(x.contract_rs) if x.contract_rs is not None else None,"from":x.effective_from.isoformat() if x.effective_from else None} for x in contracts],
                "pan":[{"pan":x.pan,"status":x.pan_status,"tds_rate":float(x.validated_tds_rate) if x.validated_tds_rate is not None else None} for x in pans],
                "compliance":[{"status":x.status,"reason":x.reason,"book_id":x.book_id,"show_id":x.show_id} for x in comps]}
    finally:db.close()
