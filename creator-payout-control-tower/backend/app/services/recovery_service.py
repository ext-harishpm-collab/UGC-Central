import uuid,re
from decimal import Decimal
from openpyxl import load_workbook
from ..db.session import SessionLocal
from ..models import RecoveryEvent,Adjustment,ContentItem
def s(v):return str(v).strip() if v not in (None,"") else None
def n(v):
    try:return Decimal(str(v))
    except:return Decimal("0")
def ingest_recovery_and_adjustments(path,period):
    wb=load_workbook(path,read_only=True,data_only=True);db=SessionLocal()
    out={"recovery_rows":0,"adjustments":0,"unmapped_adjustments":0}
    try:
        for sn in wb.sheetnames:
            low=sn.lower();ws=wb[sn]
            if low=="recovery":
                rows=list(ws.iter_rows(min_row=1,max_row=2,values_only=True))
                if len(rows)<2:continue
                hdr=[s(x) or "" for x in rows[1]]
                ia=next((i for i,x in enumerate(hdr) if "author id" in x.lower()),None)
                iopen=next((i for i,x in enumerate(hdr) if "still pending" in x.lower()),None)
                target=next((i for i,x in enumerate(hdr) if period.split()[0] in x and "paid" in x.lower()),None)
                if ia is None:continue
                for rn,row in enumerate(ws.iter_rows(min_row=3,values_only=True),3):
                    aid=s(row[ia])
                    if not aid:continue
                    opening=n(row[iopen]) if iopen is not None and iopen<len(row) else Decimal("0")
                    applied=n(row[target]) if target is not None and target<len(row) else Decimal("0")
                    if applied or opening:
                        closing=max(Decimal("0"),opening-applied)
                        db.add(RecoveryEvent(id=uuid.uuid4().hex,payment_period=period,author_id=aid,opening_outstanding=opening,applied=applied,closing_outstanding=closing,source_file=path.rsplit("/",1)[-1],source_sheet=sn,source_row=rn));out["recovery_rows"]+=1
            elif low=="adjustment":
                for rn,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
                    if len(row)<7:continue
                    book=s(row[2]);amount=n(row[3]);reason=s(row[4]);currency=s(row[5]);reward=s(row[6])
                    if not book or not amount:continue
                    content=db.query(ContentItem).filter(ContentItem.book_id==book).first()
                    db.add(Adjustment(id=uuid.uuid4().hex,payment_period=period,author_id=content.author_id if content else None,book_id=book,show_id=content.show_id if content else None,amount=amount,reason=reason or "REASON REQUIRED",currency=currency,reward_type=reward, status="PENDING"));out["adjustments"]+=1
                    out["unmapped_adjustments"]+=1 if not content else 0
        db.commit();return out
    except Exception:db.rollback();raise
    finally:db.close()
def author_recovery(period):
    db=SessionLocal()
    try:
        rows=db.query(RecoveryEvent).filter(RecoveryEvent.payment_period==period).all();out={}
        for r in rows:out[r.author_id]=out.get(r.author_id,Decimal("0"))+r.applied
        return {k:float(v) for k,v in out.items()}
    finally:db.close()
def author_adjustments(period):
    db=SessionLocal()
    try:
        rows=db.query(Adjustment).filter(Adjustment.payment_period==period,Adjustment.status=="APPROVED").all();out={}
        for r in rows:
            if r.author_id:out[r.author_id]=out.get(r.author_id,Decimal("0"))+r.amount
        return {k:float(v) for k,v in out.items()}
    finally:db.close()
