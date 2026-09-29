from pathlib import Path
import tempfile, shutil
from openpyxl import load_workbook
from ..db.session import SessionLocal
from ..models import PaymentTransaction, ContentItem, BankAccount, ReconciliationEvent
from .status import normalize_payment_status

FIXED = [
    "Payout Type","Book ID","User Bank Account","Amount Before Tax","Currency",
    "Amount After Tax","Payout Status","UTR","TDS%","Transaction Date","Payout Id","Mode","Status Details"
]

def _norm(v):
    return str(v).strip() if v not in (None, "") else None

def _money(v):
    try:
        return float(v)
    except Exception:
        return None

def _bank_account_from_json(v):
    if not isinstance(v, str):
        return _norm(v), None, None
    try:
        import json
        d=json.loads(v)
        return _norm(d.get("account_number")), _norm(d.get("ifsc")), _norm(d.get("bank_name"))
    except Exception:
        return _norm(v), None, None

def parse_uwt_workbook(path: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    rows=[]
    for ws in wb.worksheets:
        headers=None
        for r in ws.iter_rows(min_row=1, max_row=min(10, ws.max_row or 1), values_only=True):
            vals=[_norm(v) for v in r]
            if vals[:2] == ["Payout Type","Book ID"]:
                headers=vals
                break
        if not headers:
            continue
        idx={h:i for i,h in enumerate(headers) if h}
        for rn,r in enumerate(ws.iter_rows(min_row=2, values_only=True),2):
            if not any(v not in (None,"") for v in r): continue
            def val(k):
                i=idx.get(k)
                return r[i] if i is not None and i<len(r) else None
            raw_status=_norm(val("Payout Status"))
            account,ifsc,bank=_bank_account_from_json(val("User Bank Account"))
            rows.append({
                "sheet":ws.title,"row":rn,"payout_type":_norm(val("Payout Type")),
                "book_id":_norm(val("Book ID")),"account_number":account,"ifsc":ifsc,"bank_name":bank,
                "amount_before_tax":_money(val("Amount Before Tax")),"currency":_norm(val("Currency")) or "INR",
                "amount_after_tax":_money(val("Amount After Tax")),"raw_status":raw_status,
                "status":normalize_payment_status(raw_status).value,"utr":_norm(val("UTR")),
                "tds_rate":_money(val("TDS%")),"transaction_date":val("Transaction Date"),
                "payout_id":_norm(val("Payout Id")),"mode":_norm(val("Mode")),
                "status_details":_norm(val("Status Details"))
            })
    return rows

def import_uwt_workbook(path: str):
    items=parse_uwt_workbook(path)
    db=SessionLocal()
    result={"rows":len(items),"matched":0,"unmatched":0,"success":0,"failed":0,"reversed":0,"bank_evidence":0,"details":[]}
    try:
        for item in items:
            row=None
            match="NONE"
            if item["payout_id"]:
                row=db.query(PaymentTransaction).filter(PaymentTransaction.payout_id==item["payout_id"]).first()
                if row: match="PAYOUT_ID"
            if not row and item["utr"]:
                row=db.query(PaymentTransaction).filter(PaymentTransaction.utr==item["utr"]).first()
                if row: match="UTR"
            if not row and item["book_id"] and item["amount_after_tax"] is not None:
                row=db.query(PaymentTransaction).filter(
                    PaymentTransaction.book_id==item["book_id"],
                    PaymentTransaction.amount_after_tax==item["amount_after_tax"]
                ).first()
                if row: match="BOOK_AMOUNT"
            if not row:
                result["unmatched"]+=1
                result["details"].append({**item,"matched":False,"match_type":match,"reason":"No deterministic match"})
                continue

            row.status=item["status"]
            row.original_status=item["raw_status"]
            row.utr=item["utr"] or row.utr
            row.payout_id=item["payout_id"] or row.payout_id
            row.account_number=item["account_number"] or row.account_number
            row.ifsc=item["ifsc"] or row.ifsc
            row.bank_name=item["bank_name"] or row.bank_name
            row.ledger_state = "CLOSED" if item["status"]=="SUCCESS" else "RETRYABLE" if item["status"] in {"FAILED","REVERSED"} else "OPEN"

            # Only SUCCESS is trusted as successful-payment evidence.
            if item["status"]=="SUCCESS" and row.author_id and item["account_number"]:
                bank = db.query(BankAccount).filter(
                    BankAccount.author_id==row.author_id,
                    BankAccount.account_number==str(item["account_number"]),
                    BankAccount.ifsc==item["ifsc"]
                ).first()
                if not bank:
                    bank=BankAccount(
                        id=f"bank_{row.author_id}_{item['account_number']}"[:64],
                        author_id=row.author_id,
                        account_number=str(item["account_number"]),
                        ifsc=item["ifsc"], bank_name=item["bank_name"],
                        verification_status="SUCCESS_PAYMENT_EVIDENCE",
                        successful_payment_count=1, is_current=True
                    )
                    db.add(bank); result["bank_evidence"]+=1
                else:
                    bank.successful_payment_count=(bank.successful_payment_count or 0)+1
                    bank.is_current=True
                    bank.verification_status="SUCCESS_PAYMENT_EVIDENCE"

            db.add(ReconciliationEvent(
                payment_transaction_id=row.id,
                match_type=f"UWT_{match}",
                finance_status=item["status"],
                ledger_state=row.ledger_state,
                notes=item["status_details"]
            ))
            result["matched"]+=1
            result[item["status"].lower()]=result.get(item["status"].lower(),0)+1
            result["details"].append({
                "matched":True,"payment_transaction_id":row.id,"author_id":row.author_id,
                "book_id":row.book_id,"show_id":row.show_id,"status":row.status,
                "ledger_state":row.ledger_state,"match_type":match,"utr":row.utr
            })
        db.commit()
        return result
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

def fixed_uwt_rows(period=None,status=None):
    db=SessionLocal()
    try:
        q=db.query(PaymentTransaction)
        if period:q=q.filter(PaymentTransaction.payment_period==period)
        if status:q=q.filter(PaymentTransaction.status==status)
        out=[]
        for r in q.order_by(PaymentTransaction.source_row.asc()).all():
            out.append({
                "Payout Type":r.payment_type or "",
                "Book ID":r.book_id or "",
                "User Bank Account":r.account_number or "",
                "Amount Before Tax":float(r.amount_before_tax) if r.amount_before_tax is not None else "",
                "Currency":r.currency or "INR",
                "Amount After Tax":float(r.amount_after_tax) if r.amount_after_tax is not None else "",
                "Payout Status":r.status,
                "UTR":r.utr or "",
                "TDS%":"",
                "Transaction Date":r.transaction_date or "",
                "Payout Id":r.payout_id or "",
                "Mode":"",
                "Status Details":r.original_status or ""
            })
        return out
    finally:
        db.close()
