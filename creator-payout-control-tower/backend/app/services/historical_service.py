import uuid
from pathlib import Path
from openpyxl import load_workbook
from ..db.session import SessionLocal
from ..models import Author, BankAccount, ContentItem, PaymentTransaction
from .status import normalize_payment_status

def norm(v):
    return str(v).strip() if v not in (None,"") else None

def content_type(v):
    x=(v or "").upper()
    if "N2A" in x or "A2A" in x:return "N2A/A2A"
    if "SERIES" in x:return "SERIES"
    if "NOVEL" in x:return "NOVEL"
    return None

def headers(ws):
    return [norm(v) or f"column_{i+1}" for i,v in enumerate(next(ws.iter_rows(min_row=1,max_row=1,values_only=True)))]

def idx(hs,*names):
    for i,h in enumerate(hs):
        x=h.lower()
        if any(n.lower() in x for n in names):return i
    return None

def backfill_files(paths:list[str])->dict:
    db=SessionLocal();result={"files":0,"authors":0,"content":0,"successful_payments":0,"payment_events":0}
    try:
        book_show={}
        book_type={}
        staged=[]
        for path in paths:
            wb=load_workbook(path,read_only=True,data_only=True);result["files"]+=1
            for s in wb.sheetnames:
                ws=wb[s]; sn=s.lower()
                hs=headers(ws)
                if "show book mapping" in sn:
                    ib=idx(hs,"book id");isid=idx(hs,"show id");ia=idx(hs,"author uid")
                    for row in ws.iter_rows(min_row=2,values_only=True):
                        if ib is None:continue
                        b=norm(row[ib]); sh=norm(row[isid]) if isid is not None and isid<len(row) else None; a=norm(row[ia]) if ia is not None and ia<len(row) else None
                        if b:book_show[b]=(sh,a)
                elif s in {"RS Final","Inc Final"} or "rs final" in sn or "inc final" in sn:
                    ib=idx(hs,"book id");it=idx(hs,"ip type")
                    if ib is not None:
                        for row in ws.iter_rows(min_row=3,values_only=True):
                            b=norm(row[ib]);t=content_type(norm(row[it]) if it is not None and it<len(row) else None)
                            if b and t:book_type[b]=t
                elif "utr details" in sn:
                    for row in ws.iter_rows(min_row=2,values_only=True):
                        if len(row)<10:continue
                        status=normalize_payment_status(norm(row[8]))
                        if status.value=="UNMATCHED":continue
                        staged.append({"author_id":None,"book_id":None,"show_id":None,"status":status.value,"utr":norm(row[9]),
                                       "account_number":norm(row[2]),"ifsc":norm(row[3]),"amount_after_tax":row[7],"source_sheet":s})
                elif "payment status dump" in sn:
                    for row in ws.iter_rows(min_row=2,values_only=True):
                        if len(row)<13:continue
                        bid=norm(row[2]);st=normalize_payment_status(norm(row[3])).value
                        staged.append({"author_id":None,"book_id":bid,"show_id":None,"status":st,"utr":norm(row[11]),
                                       "amount_after_tax":row[10],"source_sheet":s})
                elif sn=="author level novel" or sn=="author level series" or sn=="author level n2aa2a":
                    ia=idx(hs,"author id");iname=idx(hs,"author name");iban=idx(hs,"acc no");iifsc=idx(hs,"ifsc code")
                    if ia is not None:
                        for row in ws.iter_rows(min_row=3,values_only=True):
                            aid=norm(row[ia])
                            if not aid:continue
                            a=db.query(Author).filter_by(author_id=aid).first()
                            if not a:
                                a=Author(id=uuid.uuid4().hex,author_id=aid,name=norm(row[iname]) if iname is not None else None,status="HISTORICAL");db.add(a);result["authors"]+=1
                            # Bank data here is reference only; do not mark as successful-payment evidence.
        for e in staged:
            if e["book_id"] in book_show:
                e["show_id"]=book_show[e["book_id"]][0]
                e["author_id"]=e["author_id"] or book_show[e["book_id"]][1]
            if e["book_id"] in book_type:e["content_type"]=book_type[e["book_id"]]
            # A successful bank row is evidence only when the finance result is SUCCESS.
            tx=PaymentTransaction(id=uuid.uuid4().hex,author_id=e.get("author_id"),book_id=e.get("book_id"),show_id=e.get("show_id"),
                                  payment_type="HISTORICAL_RESULT",amount_after_tax=e.get("amount_after_tax"),currency="INR",
                                  status=e["status"],utr=e.get("utr"),original_status=e["status"],source_sheet=e["source_sheet"],
                                  payment_period=None,content_type=e.get("content_type"),ledger_state="CLOSED" if e["status"]=="SUCCESS" else "RETRYABLE")
            db.add(tx);result["payment_events"]+=1
            if e["status"]=="SUCCESS":result["successful_payments"]+=1
        for book,(show,author) in book_show.items():
            t=book_type.get(book)
            q=db.query(ContentItem).filter(ContentItem.book_id==book)
            item=q.first()
            if not item:
                db.add(ContentItem(id=uuid.uuid4().hex,book_id=book,show_id=show,author_id=author,content_type=t));result["content"]+=1
        db.commit();return result
    except Exception:
        db.rollback();raise
    finally:db.close()
