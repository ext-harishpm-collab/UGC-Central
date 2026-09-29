import csv, io, uuid
from collections import defaultdict
from decimal import Decimal
from openpyxl import load_workbook

def s(v): return str(v).strip() if v not in (None,"") else ""
def d(v):
    try:return Decimal(str(v))
    except:return Decimal("0")

def find_header(ws, max_rows=10):
    best=(0,None,None)
    for rn,row in enumerate(ws.iter_rows(min_row=1,max_row=min(max_rows,ws.max_row or max_rows),values_only=True),1):
        vals=[s(x).lower() for x in row]; hits=sum(1 for x in vals if x in {"book id","show id","author id","author uid","pay?","ip type","gross","net amount"})
        if hits>best[0]:best=(hits,rn,[s(x) for x in row])
    return best[1],best[2]

def idx(headers,*names):
    for i,h in enumerate(headers):
        x=h.lower()
        if any(n.lower() in x for n in names): return i
    return None

def read_dump(path, source_kind):
    wb=load_workbook(path,read_only=True,data_only=True)
    rows=[]
    for ws in wb.worksheets:
        hr,headers=find_header(ws)
        if not hr:continue
        m={h:i for i,h in enumerate(headers) if h}
        for rn,row in enumerate(ws.iter_rows(min_row=hr+1,values_only=True),hr+1):
            if not any(v not in (None,"") for v in row):continue
            bid= s(row[m["Book ID"]]) if "Book ID" in m and m["Book ID"]<len(row) else None
            show= s(row[m["Show ID"]]) if "Show ID" in m and m["Show ID"]<len(row) else None
            author=s(row[m["Author ID"]]) if "Author ID" in m and m["Author ID"]<len(row) else s(row[m["Author UID"]]) if "Author UID" in m and m["Author UID"]<len(row) else None
            gross=row[m["Gross"]] if "Gross" in m and m["Gross"]<len(row) else 0
            tds=row[m["TDS Amount"]] if "TDS Amount" in m and m["TDS Amount"]<len(row) else 0
            net=row[m["Net Amount"]] if "Net Amount" in m and m["Net Amount"]<len(row) else d(gross)-d(tds)
            ip=s(row[m["IP Type"]]) if "IP Type" in m and m["IP Type"]<len(row) else ""
            pay=s(row[m["Pay?"]]) if "Pay?" in m and m["Pay?"]<len(row) else ""
            rows.append({"source_kind":source_kind,"source_file":path.rsplit("/",1)[-1],"source_sheet":ws.title,"source_row":rn,
                         "book_id":bid,"show_id":show,"author_id":author,"ip_type":ip,"pay_flag":pay,
                         "gross":str(d(gross)),"tds":str(d(tds)),"net":str(d(net))})
    return rows

def process_two_dumps(inc_path, rs_path, period):
    inc=read_dump(inc_path,"INCENTIVE_DUMP")
    rs=read_dump(rs_path,"REVENUE_SHARE_DUMP")
    all_rows=inc+rs
    grouped=defaultdict(lambda:{"inc":Decimal("0"),"rs":Decimal("0"),"rows":[],"types":set(),"source_rows":0})
    qc=[]
    for r in all_rows:
        key=(r["author_id"],r["book_id"],r["show_id"])
        g=grouped[key];g["rows"].append(r);g["source_rows"]+=1
        if r["ip_type"]:g["types"].add(r["ip_type"])
        eligible=(r["pay_flag"].lower() in {"yes","eligible","paid",""})
        if not eligible:
            qc.append({"severity":"REVIEW","rule_id":"PAY-FLAG","message":"Source Pay? is not eligible","author_id":r["author_id"],"book_id":r["book_id"],"show_id":r["show_id"],"source_sheet":r["source_sheet"],"source_row":r["source_row"]})
            continue
        typ=(r["ip_type"] or "").upper()
        if r["source_kind"]=="INCENTIVE_DUMP":
            if "N2A" in typ or "A2A" in typ:
                qc.append({"severity":"BLOCK","rule_id":"N2A-INC","message":"N2A/A2A incentive excluded","author_id":r["author_id"],"book_id":r["book_id"],"show_id":r["show_id"],"source_sheet":r["source_sheet"],"source_row":r["source_row"]})
            else:g["inc"]+=d(r["gross"])
        else:
            g["rs"]+=d(r["gross"])
    output=[];author=defaultdict(lambda:{"inc":Decimal("0"),"rs":Decimal("0"),"gross":Decimal("0"),"tds":Decimal("0"),"net":Decimal("0"),"rows":0,"lineage":[]})
    for (aid,bid,sid),g in grouped.items():
        inc=g["inc"];rs=g["rs"];gross=inc+rs; 
        for r in g["rows"]:
            if r["source_kind"]=="INCENTIVE_DUMP" and not ("N2A" in r["ip_type"].upper() or "A2A" in r["ip_type"].upper()):
                author[aid]["inc"]+=d(r["gross"])
            if r["source_kind"]=="REVENUE_SHARE_DUMP":
                author[aid]["rs"]+=d(r["gross"])
            author[aid]["tds"]+=d(r["tds"]);author[aid]["net"]+=d(r["net"]);author[aid]["rows"]+=1;author[aid]["lineage"].append(r)
        output.append({"period":period,"author_id":aid,"book_id":bid,"show_id":sid,"incentive_gross":float(inc),"revenue_share_gross":float(rs),"gross":float(gross),
                       "source_row_count":g["source_rows"],"sources":"; ".join(sorted({r["source_kind"]+":"+r["source_sheet"]+":"+str(r["source_row"]) for r in g["rows"]}))})
    author_rows=[]
    for aid,g in author.items():
        author_rows.append({"period":period,"author_id":aid,"incentive_gross":float(g["inc"]),"revenue_share_gross":float(g["rs"]),"gross":float(g["inc"]+g["rs"]),"source_row_count":g["rows"],
                            "source_refs":" | ".join(f'{r["source_kind"]}:{r["source_file"]}:{r["source_sheet"]}:{r["source_row"]}' for r in g["lineage"])})
    return {"period":period,"input_summary":{"incentive_rows":len(inc),"revenue_share_rows":len(rs),"total_source_rows":len(all_rows)},
            "show_level_rows":output,"author_level_rows":author_rows,"qc_rows":qc,
            "qc_summary":{"BLOCK":sum(x["severity"]=="BLOCK" for x in qc),"REVIEW":sum(x["severity"]=="REVIEW" for x in qc),
                          "PASS":len(output)-sum(x["severity"]=="BLOCK" for x in qc)}}

def csv_bytes(rows, headers):
    buf=io.StringIO();w=csv.DictWriter(buf,fieldnames=headers,extrasaction="ignore");w.writeheader();w.writerows(rows);return buf.getvalue().encode("utf-8-sig")
