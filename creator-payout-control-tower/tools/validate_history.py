import argparse, json
from collections import Counter, defaultdict
from openpyxl import load_workbook

def s(v): return str(v).strip() if v not in (None,"") else ""

def validate(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    out={"file":path.rsplit("/",1)[-1],"sheets":len(wb.sheetnames),"uwt":None,"utr":None,"earnings":{},"mapping":{}}
    if "uwt final" in wb.sheetnames:
        ws=wb["uwt final"]; c=Counter()
        rows=0
        for r in ws.iter_rows(min_row=2,values_only=True):
            if any(v not in (None,"") for v in r):
                rows+=1
                if len(r)>6 and r[6] not in (None,""): c[s(r[6]).lower()]+=1
        out["uwt"]={"rows":rows,"status_counts":dict(c)}
    if "UTR Details" in wb.sheetnames:
        ws=wb["UTR Details"]; c=Counter();rows=0
        for r in ws.iter_rows(min_row=2,values_only=True):
            if any(v not in (None,"") for v in r):
                rows+=1
                if len(r)>8 and r[8] not in (None,""): c[s(r[8]).lower()]+=1
        out["utr"]={"rows":rows,"status_counts":dict(c)}
    for name in ["Inc Final","RS Final"]:
        if name not in wb.sheetnames: continue
        ws=wb[name]
        hdr=[s(v) for v in next(ws.iter_rows(min_row=2,max_row=2,values_only=True))]
        idx={h.lower():i for i,h in enumerate(hdr)}
        ip=idx.get("ip type"); pay=idx.get("pay?"); book=idx.get("book id")
        count=0; pay_counts=Counter(); types=Counter(); missing_book=0
        for r in ws.iter_rows(min_row=3,values_only=True):
            if not any(v not in (None,"") for v in r): continue
            count+=1
            if pay is not None: pay_counts[s(r[pay]).lower()]+=1
            if ip is not None: types[s(r[ip]) or "BLANK"]+=1
            if book is None or book>=len(r) or not s(r[book]): missing_book+=1
        out["earnings"][name]={"rows":count,"pay_flags":dict(pay_counts),"ip_types":dict(types),"missing_book_id":missing_book}
    if "Show Book Mapping" in wb.sheetnames:
        ws=wb["Show Book Mapping"]; hdr=[s(v) for v in next(ws.iter_rows(min_row=1,max_row=1,values_only=True))]
        idx={h.lower():i for i,h in enumerate(hdr)}; ib=idx.get("book id"); ish=idx.get("show id");ia=idx.get("author uid")
        books=defaultdict(set); shows=defaultdict(set)
        for r in ws.iter_rows(min_row=2,values_only=True):
            if ib is None or ib>=len(r):continue
            b=s(r[ib]); sh=s(r[ish]) if ish is not None and ish<len(r) else ""; a=s(r[ia]) if ia is not None and ia<len(r) else ""
            if b:books[b].add(sh or "MISSING_SHOW")
            if sh:shows[sh].add(b or "MISSING_BOOK")
        multi_book=sum(1 for v in books.values() if len(v)>1)
        multi_show=sum(1 for v in shows.values() if len(v)>1)
        out["mapping"]={"book_ids":len(books),"show_ids":len(shows),"book_to_multiple_shows":multi_book,"show_to_multiple_books":multi_show}
    return out

if __name__=="__main__":
    ap=argparse.ArgumentParser();ap.add_argument("files",nargs="+");args=ap.parse_args()
    print(json.dumps([validate(p) for p in args.files],indent=2,default=str))
